"""Excel parser (openpyxl): every sheet, merged cells unmerged, empty rows split tables (§10.3)."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.ingestion.ir import IRBuilder, cell_text

MAX_CELLS = 100_000


def parse_xlsx(path: Path, ir: IRBuilder) -> None:
    workbook = load_workbook(path, data_only=True)  # formulas -> cached values
    cells_seen = 0
    for sheet in workbook.worksheets:
        ir.count("sheets")
        if sheet.sheet_state != "visible":
            ir.warnings.append(f"sheet {sheet.title!r} is hidden; included")
        max_row, max_col = sheet.max_row, sheet.max_column
        cells_seen += max_row * max_col
        if cells_seen > MAX_CELLS:
            ir.warnings.append(f"stopped at sheet {sheet.title!r}: more than {MAX_CELLS} cells")
            break
        grid = [
            [cell_text(sheet.cell(r, c).value) for c in range(1, max_col + 1)]
            for r in range(1, max_row + 1)
        ]
        # Unmerge: copy the top-left value into every covered cell so each row stands alone.
        for merged in sheet.merged_cells.ranges:
            value = grid[merged.min_row - 1][merged.min_col - 1]
            for r in range(merged.min_row, merged.max_row + 1):
                for c in range(merged.min_col, merged.max_col + 1):
                    grid[r - 1][c - 1] = value
        _split_blocks(grid, sheet, ir)


def _split_blocks(grid: list[list[str]], sheet: Worksheet, ir: IRBuilder) -> None:
    """Runs of non-empty rows become blocks; a run that is a single cell becomes text."""
    start: int | None = None
    for index in range(len(grid) + 1):
        empty = index == len(grid) or not any(grid[index])
        if not empty and start is None:
            start = index
        elif empty and start is not None:
            _emit(grid[start:index], start, sheet, ir)
            start = None


def _emit(rows: list[list[str]], first_row: int, sheet: Worksheet, ir: IRBuilder) -> None:
    filled = [(r, c) for r, row in enumerate(rows) for c, v in enumerate(row) if v]
    cols = sorted({c for _, c in filled})
    ref_range = (
        f"{get_column_letter(cols[0] + 1)}{first_row + 1}:"
        f"{get_column_letter(cols[-1] + 1)}{first_row + len(rows)}"
    )
    ref = {"sheet": sheet.title, "range": ref_range}
    if len(rows) == 1 and len(set(rows[0]) - {""}) == 1:
        bold = bool(sheet.cell(first_row + 1, cols[0] + 1).font.bold)
        ir.text("heading" if bold else "paragraph", rows[0][cols[0]], ref)
        return
    ir.table(rows, ref)
