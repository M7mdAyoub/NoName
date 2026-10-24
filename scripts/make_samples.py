"""Write the synthetic sample framework in every supported format to tests/fixtures/.

    python scripts/make_samples.py

The content is invented (no real project data): a small agri-food logframe with an impact,
two outcomes and three outputs, indicators with baselines, targets, sources and frequencies.
"""

from __future__ import annotations

import csv
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "tests" / "fixtures"

TITLE = "Results Framework - Green Valley Livelihoods Programme (synthetic)"
HEADER = ["Level", "Code", "Result statement", "Indicator", "Unit", "Baseline", "Target",
          "Means of verification", "Frequency"]  # fmt: skip
ROWS = [
    ["Impact", "0", "Household incomes in the valley increased", "Average monthly household income", "JOD", "310", "400", "Household survey", "Annual"],
    ["Outcome", "1", "Women's participation in agri-food production increased", "Number of women completing agri-food training", "people", "0", "400", "Training attendance records", "Quarterly"],
    ["Outcome", "1", "Women's participation in agri-food production increased", "% of trained women applying new practices", "%", "0", "60", "Follow-up survey", "Semiannual"],
    ["Output", "1.1", "Training courses delivered", "Number of training sessions held", "sessions", "0", "80", "Session reports", "Quarterly"],
    ["Output", "1.2", "Start-up kits distributed", "Number of start-up kits distributed", "kits", "0", "250", "Distribution lists", "Quarterly"],
    ["Outcome", "2", "Market access for small producers improved", "Total sales through cooperative markets", "JOD", "12000", "50000", "Market sales records", "Quarterly"],
    ["Output", "2.1", "Cooperative market platform launched", "Market platform operational", "milestone", "No", "Yes", "Project reports", "Once"],
]  # fmt: skip
ASSUMPTIONS = [
    ["Assumption", "Risk level"],
    ["Security situation remains stable", "Medium"],
    ["Municipal partners keep co-funding", "Low"],
]


def write_xlsx(path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Logframe"
    ws["A1"] = TITLE
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:I1")
    ws.append([])
    ws.append(HEADER)
    for row in ROWS:
        ws.append([float(v) if v.replace(".", "").isdigit() else v for v in row])
    # Real logframes merge the level/statement cells of an outcome with several indicators.
    ws.merge_cells("A5:A6")
    ws.merge_cells("B5:B6")
    ws.merge_cells("C5:C6")  # openpyxl keeps only the top-left value, as Excel does
    ws.append([])
    for row in ASSUMPTIONS:
        ws.append(row)
    notes = wb.create_sheet("Notes")
    notes["A1"] = "Targets are end-of-programme values."
    notes.sheet_state = "hidden"
    wb.save(path)


def write_csv(path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(HEADER)
        writer.writerows(ROWS)


def write_docx(path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading(TITLE, level=1)
    doc.add_paragraph(
        "This monitoring plan describes how progress of the programme is measured. "
        "Indicators are reported quarterly unless stated otherwise."
    )
    doc.add_heading("Indicators", level=2)
    table = doc.add_table(rows=1, cols=len(HEADER))
    for cell, title in zip(table.rows[0].cells, HEADER, strict=True):
        cell.text = title
    for row in ROWS:
        cells = table.add_row().cells
        for cell, value in zip(cells, row, strict=True):
            cell.text = value
    doc.add_heading("Assumptions", level=2)
    for assumption, _ in ASSUMPTIONS[1:]:
        doc.add_paragraph(assumption, style="List Bullet")
    doc.save(str(path))


def _pdf_escape(value: str) -> str:
    return value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def write_pdf(path: Path) -> None:
    """A one-page digital PDF: a title, a paragraph and a ruled table (pdfplumber finds it)."""
    columns = ["Code", "Indicator", "Baseline", "Target"]
    data = [[r[1], r[3][:48], r[5], r[6]] for r in ROWS]
    widths = [40, 300, 70, 70]
    x0, top, row_h = 40, 700, 18
    ops = [
        f"BT /F1 13 Tf 40 760 Td ({_pdf_escape(TITLE)}) Tj ET",
        "BT /F1 10 Tf 40 740 Td (Indicator targets for the programme, by result code.) Tj ET",
        "0.5 w",
    ]
    table = [columns, *data]
    total_w = sum(widths)
    for i in range(len(table) + 1):  # horizontal rules
        y = top - i * row_h
        ops.append(f"{x0} {y} m {x0 + total_w} {y} l S")
    x = x0
    for w in [0, *widths]:  # vertical rules
        x += w
        ops.append(f"{x} {top} m {x} {top - len(table) * row_h} l S")
    for i, row in enumerate(table):
        x = x0
        for w, value in zip(widths, row, strict=True):
            y = top - (i + 1) * row_h + 5
            ops.append(f"BT /F1 8 Tf {x + 3} {y} Td ({_pdf_escape(value)}) Tj ET")
            x += w
    stream = "\n".join(ops).encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    path.write_bytes(bytes(out))


def main() -> int:
    sys.path.insert(0, str(ROOT))
    OUT.mkdir(parents=True, exist_ok=True)
    for suffix, writer in (
        ("xlsx", write_xlsx),
        ("csv", write_csv),
        ("docx", write_docx),
        ("pdf", write_pdf),
    ):
        target = OUT / f"sample_logframe.{suffix}"
        writer(target)
        print(f"wrote {target.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
